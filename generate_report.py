"""
Script para generar el documento Word de la Práctica 2.1
Notificaciones Push y Acceso a Características Nativas en PWA
"""
import os
import glob
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

# Directorio de capturas de pantalla
ARTIFACTS_DIR = r"C:\Users\Angel Medina\.gemini\antigravity-ide\brain\935c7644-9f99-4c8c-9578-ee0714a22083"
OUTPUT_PATH = r"C:\Users\Angel Medina\Desktop\files\Practica_2_1_PWA_Notificaciones_Push.docx"

# ── Helpers ────────────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

def add_heading(doc, text, level=1, color=None):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if color:
        for run in p.runs:
            run.font.color.rgb = RGBColor(*color)
    return p

def add_paragraph(doc, text, bold=False, italic=False, color=None, size=None, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = alignment
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    if size:
        run.font.size = Pt(size)
    return p

def add_code_block(doc, code, title=""):
    if title:
        p = doc.add_paragraph()
        run = p.add_run(f"  {title}")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(14, 165, 233)
    
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(code)
    run.font.name = 'Courier New'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(34, 197, 94)
    
    # Fondo oscuro para el bloque de código
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), '1E293B')
    shd.set(qn('w:val'), 'clear')
    pPr.append(shd)
    return p

def add_screenshot(doc, img_path, caption, width_inches=6.0):
    if not os.path.exists(img_path):
        print(f"[WARN] Imagen no encontrada: {img_path}")
        p = doc.add_paragraph()
        run = p.add_run(f"[Imagen no disponible: {os.path.basename(img_path)}]")
        run.italic = True
        run.font.color.rgb = RGBColor(156, 163, 175)
        return
    try:
        doc.add_picture(img_path, width=Inches(width_inches))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(f"Figura: {caption}")
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(107, 114, 128)
        print(f"[OK] Imagen: {caption}")
    except Exception as e:
        print(f"[ERROR] Falló imagen {img_path}: {e}")

def find_screenshot(name):
    """Busca un screenshot por nombre en el directorio de artefactos."""
    pattern = os.path.join(ARTIFACTS_DIR, f"{name}_*.png")
    matches = glob.glob(pattern)
    if matches:
        return sorted(matches)[-1]  # Más reciente
    # Intentar sin timestamp
    direct = os.path.join(ARTIFACTS_DIR, f"{name}.png")
    if os.path.exists(direct):
        return direct
    return None

# ── Crear documento ────────────────────────────────────────────────────────────
def create_word_document():
    doc = Document()
    
    # Configurar márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3)
        section.right_margin  = Cm(2.5)
    
    # ── PORTADA ────────────────────────────────────────────────────────────────
    doc.add_paragraph()
    doc.add_paragraph()
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PRÁCTICA 2.1")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(14, 165, 233)
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Notificaciones Push y Acceso a Características Nativas en PWA")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph()
    
    app_name = doc.add_paragraph()
    app_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = app_name.add_run("Aplicación: HealthVault — Sistema Médico de Gestión")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(71, 85, 105)
    run.italic = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Info del alumno
    table_info = doc.add_table(rows=3, cols=2)
    table_info.style = 'Table Grid'
    labels = ["Materia:", "Fecha:", "Tecnologías:"]
    values = [
        "Desarrollo de Aplicaciones Web Móviles",
        datetime.datetime.now().strftime("%d de %B de %Y"),
        "Flask (Python) | Service Worker | Push API | Geolocation API | MediaDevices API"
    ]
    for i, (label, val) in enumerate(zip(labels, values)):
        row = table_info.rows[i]
        set_cell_bg(row.cells[0], '0EA5E9')
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        row.cells[1].text = val
    
    doc.add_page_break()
    
    # ── SECCIÓN 1: OBJETIVO ────────────────────────────────────────────────────
    add_heading(doc, "1. Objetivo", level=1, color=(14, 165, 233))
    add_paragraph(doc,
        "Programar notificaciones push y el acceso a características nativas del dispositivo "
        "desde una Progressive Web App (PWA). Se implementó sobre la aplicación HealthVault, un "
        "sistema médico de gestión de expedientes basado en Flask/Python.")
    
    doc.add_paragraph()
    
    # ── SECCIÓN 2: MARCO TEÓRICO ───────────────────────────────────────────────
    add_heading(doc, "2. Marco Teórico — Notificaciones Push en PWA", level=1, color=(14, 165, 233))
    
    add_heading(doc, "2.1 ¿Qué son las Notificaciones Push?", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "Las notificaciones push son mensajes que un servidor envía a un dispositivo cliente, incluso "
        "cuando el usuario no tiene la aplicación web abierta. En una PWA esto se logra mediante la "
        "combinación de tres tecnologías: la Notifications API del navegador, la Push API y los "
        "Service Workers.")
    
    add_heading(doc, "2.2 Flujo de Suscripción Push (VAPID)", level=2, color=(30, 41, 59))
    
    steps_table = doc.add_table(rows=6, cols=2)
    steps_table.style = 'Table Grid'
    steps_data = [
        ("Paso 1", "El usuario carga la PWA y el Service Worker se registra."),
        ("Paso 2", "Se solicita permiso de notificaciones con Notification.requestPermission()."),
        ("Paso 3", "Si se concede, PushManager.subscribe() crea una suscripción con clave VAPID pública."),
        ("Paso 4", "El endpoint único de suscripción se envía al servidor Flask (POST /api/push/subscribe)."),
        ("Paso 5", "El servidor usa pywebpush para enviar mensajes cifrados al Push Service del navegador."),
        ("Paso 6", "El SW intercepta el evento 'push' y llama a showNotification() para mostrar la alerta."),
    ]
    for i, (paso, desc) in enumerate(steps_data):
        row = steps_table.rows[i]
        set_cell_bg(row.cells[0], '0F172A')
        row.cells[0].text = paso
        run = row.cells[0].paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(14, 165, 233)
        row.cells[1].text = desc
    
    doc.add_paragraph()
    
    # ── SECCIÓN 3: IMPLEMENTACIÓN ─────────────────────────────────────────────
    add_heading(doc, "3. Implementación", level=1, color=(14, 165, 233))
    
    add_heading(doc, "3.1 Archivos Creados/Modificados", level=2, color=(30, 41, 59))
    
    files_table = doc.add_table(rows=7, cols=3)
    files_table.style = 'Table Grid'
    files_data = [
        ("Archivo", "Tipo", "Descripción"),
        ("static/sw.js", "Service Worker", "Cache, push events, offline fallback"),
        ("static/js/pwa.js", "JavaScript", "Módulo: notificaciones, geolocalización, cámara"),
        ("static/css/pwa.css", "CSS", "Estilos premium del panel PWA"),
        ("templates/pwa_features.html", "HTML", "Página principal de la Práctica 2.1"),
        ("app.py (rutas añadidas)", "Python/Flask", "manifest.json, /pwa-features, /api/push/*"),
        (".env (actualizado)", "Config", "Claves VAPID_PUBLIC_KEY y VAPID_PRIVATE_KEY"),
    ]
    for i, row_data in enumerate(files_data):
        row = files_table.rows[i]
        for j, cell_text in enumerate(row_data):
            row.cells[j].text = cell_text
            if i == 0:
                set_cell_bg(row.cells[j], '0EA5E9')
                row.cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                row.cells[j].paragraphs[0].runs[0].bold = True
    
    doc.add_paragraph()
    
    add_heading(doc, "3.2 Service Worker (sw.js)", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "El Service Worker es el componente central de la PWA. Se registra en el navegador y "
        "funciona en segundo plano para interceptar peticiones de red (caché offline) y recibir "
        "notificaciones push del servidor.")
    
    add_code_block(doc, 
        "// Instalación: Pre-cache de recursos esenciales\n"
        "self.addEventListener('install', event => {\n"
        "  event.waitUntil(\n"
        "    caches.open(CACHE_NAME).then(cache => cache.addAll(CACHE_ASSETS))\n"
        "  );\n"
        "});\n\n"
        "// Recepción de push del servidor\n"
        "self.addEventListener('push', event => {\n"
        "  const data = event.data.json();\n"
        "  event.waitUntil(\n"
        "    self.registration.showNotification(data.title, {\n"
        "      body: data.body,\n"
        "      icon: '/static/logo.png',\n"
        "      actions: [{ action: 'open', title: 'Abrir HealthVault' }]\n"
        "    })\n"
        "  );\n"
        "});",
        title="static/sw.js — Eventos principales"
    )
    
    doc.add_paragraph()
    
    add_heading(doc, "3.3 Solicitud de Permiso de Notificaciones", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "El permiso se solicita mediante la Notifications API. Se respetan tres estados posibles: "
        "'granted' (concedido), 'denied' (bloqueado) y 'default' (pendiente). El código maneja "
        "cada caso de forma diferente.")
    
    add_code_block(doc,
        "async function requestNotificationPermission() {\n"
        "  let permission = Notification.permission;\n"
        "  if (permission === 'default') {\n"
        "    // Solicitar al usuario explícitamente\n"
        "    permission = await Notification.requestPermission();\n"
        "  }\n"
        "  if (permission === 'granted') {\n"
        "    showToast('🔔 Permisos concedidos!', 'success');\n"
        "    await subscribeToPush(); // Suscribir al push server\n"
        "  } else if (permission === 'denied') {\n"
        "    showToast('🚫 Notificaciones bloqueadas', 'warning');\n"
        "    logPermissionEvent('notifications', 'denied');\n"
        "  }\n"
        "}",
        title="static/js/pwa.js — Solicitud de permiso"
    )
    
    doc.add_paragraph()
    doc.add_page_break()
    
    # ── SECCIÓN 4: CAPTURAS DE PANTALLA ───────────────────────────────────────
    add_heading(doc, "4. Evidencias — Capturas de Pantalla", level=1, color=(14, 165, 233))
    
    # 4.1 Login
    add_heading(doc, "4.1 Página de Inicio de Sesión", level=2, color=(30, 41, 59))
    add_paragraph(doc, "Pantalla de login del sistema HealthVault. Se accede con las credenciales del administrador del sistema (admin_sys).")
    
    img = find_screenshot("01_login_page")
    add_screenshot(doc, img, "Pantalla de inicio de sesión — HealthVault PWA")
    
    doc.add_paragraph()
    
    img2 = find_screenshot("02_login_filled")
    if img2:
        add_screenshot(doc, img2, "Formulario completado con credenciales admin_sys / Admin123!")
    
    doc.add_page_break()
    
    # 4.2 Dashboard
    add_heading(doc, "4.2 Panel Principal con Enlace PWA", level=2, color=(30, 41, 59))
    add_paragraph(doc, "Tras el inicio de sesión exitoso, el dashboard muestra el enlace '🚀 PWA' en la barra de navegación, que da acceso al módulo de la Práctica 2.1.")
    
    img = find_screenshot("03_dashboard")
    add_screenshot(doc, img, "Dashboard de HealthVault mostrando el enlace 🚀 PWA en la navbar")
    
    doc.add_page_break()
    
    # 4.3 Página PWA
    add_heading(doc, "4.3 Panel PWA — Práctica 2.1", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "La página /pwa-features muestra el panel completo de la práctica. Incluye: "
        "explicación teórica del flujo de notificaciones push, tarjetas interactivas para "
        "Service Worker, Notificaciones, Geolocalización y Cámara, y el log de permisos en tiempo real.")
    
    img = find_screenshot("04_pwa_features_page")
    add_screenshot(doc, img, "Panel principal de la Práctica 2.1 — Hero banner y sección teórica")
    
    doc.add_paragraph()
    
    img = find_screenshot("05_pwa_cards")
    add_screenshot(doc, img, "Tarjetas interactivas: Service Worker, Notificaciones, Geolocalización y Cámara")
    
    doc.add_page_break()
    
    # 4.4 Notificaciones
    add_heading(doc, "4.4 Notificaciones Push — Solicitud de Permiso", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "Al hacer click en '🔔 Solicitar Permiso de Notificaciones', el navegador muestra el diálogo "
        "de permisos. Cuando el usuario acepta, se actualiza el estado en la tarjeta y se procede "
        "a suscribir al servicio push con claves VAPID.")
    
    img = find_screenshot("06_notif_permission")
    add_screenshot(doc, img, "Diálogo de permisos de notificaciones del navegador")
    
    doc.add_paragraph()
    
    img = find_screenshot("06b_notif_granted")
    add_screenshot(doc, img, "Estado tras conceder el permiso — badge '✅ Concedido'")
    
    doc.add_page_break()
    
    # 4.5 Notificación de prueba
    add_heading(doc, "4.5 Envío de Notificación Push de Prueba", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "Una vez con permiso concedido, el botón '📤 Enviar Notificación de Prueba' usa el "
        "Service Worker para mostrar una notificación del sistema. La notificación incluye título, "
        "cuerpo, ícono y acciones interactivas (botones).")
    
    img = find_screenshot("07_test_notification")
    add_screenshot(doc, img, "Toast de confirmación y notificación push enviada")
    
    doc.add_paragraph()
    
    # Código envío
    add_code_block(doc,
        "// Envío de notificación vía Service Worker\n"
        "await swRegistration.showNotification('🏥 HealthVault – Prueba', {\n"
        "  body: 'Las notificaciones push funcionan correctamente!',\n"
        "  icon: '/static/logo.png',\n"
        "  vibrate: [200, 100, 200],\n"
        "  actions: [\n"
        "    { action: 'open', title: '📂 Ir al Dashboard' },\n"
        "    { action: 'dismiss', title: '✕ Cerrar' }\n"
        "  ]\n"
        "});",
        title="Código de envío de notificación push"
    )
    
    doc.add_page_break()
    
    # 4.6 Geolocalización
    add_heading(doc, "4.6 Acceso a Característica Nativa — Geolocalización", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "Se accede a la API de Geolocalización del dispositivo mediante navigator.geolocation.getCurrentPosition(). "
        "El sistema muestra latitud, longitud, precisión, altitud y un enlace directo a Google Maps. "
        "Se implementa manejo de los tres errores posibles: permiso denegado, posición no disponible y timeout.")
    
    img = find_screenshot("08_geolocation")
    add_screenshot(doc, img, "Solicitud de permiso de geolocalización GPS")
    
    doc.add_paragraph()
    
    img = find_screenshot("08b_geo_result")
    add_screenshot(doc, img, "Coordenadas GPS obtenidas — Latitud, Longitud y Precisión")
    
    doc.add_paragraph()
    
    add_code_block(doc,
        "navigator.geolocation.getCurrentPosition(\n"
        "  (position) => {\n"
        "    const { latitude, longitude, accuracy } = position.coords;\n"
        "    // Permiso concedido: mostrar coordenadas\n"
        "    updatePermissionBadge('geo-badge', 'granted');\n"
        "    showCoordinates(latitude, longitude, accuracy);\n"
        "  },\n"
        "  (error) => {\n"
        "    // Permiso denegado u otro error\n"
        "    if (error.code === error.PERMISSION_DENIED) {\n"
        "      showToast('🚫 Permiso denegado', 'error');\n"
        "    }\n"
        "  },\n"
        "  { enableHighAccuracy: true, timeout: 10000 }\n"
        ");",
        title="Código de acceso a geolocalización con manejo de permisos"
    )
    
    doc.add_page_break()
    
    # 4.7 Cámara
    add_heading(doc, "4.7 Acceso a Cámara y Micrófono (MediaDevices API)", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "Se accede a la cámara y micrófono del dispositivo mediante navigator.mediaDevices.getUserMedia(). "
        "La página ofrece tres modos: Cámara+Micrófono, Solo Cámara y Solo Micrófono. Cuando el permiso "
        "es concedido, se activa el preview de video en tiempo real con opción de capturar fotos.")
    
    img = find_screenshot("09_camera_permission")
    add_screenshot(doc, img, "Diálogo de permiso de cámara y micrófono")
    
    doc.add_paragraph()
    
    img = find_screenshot("09b_camera_active")
    add_screenshot(doc, img, "Cámara activa — Preview de video en tiempo real")
    
    doc.add_paragraph()
    
    add_code_block(doc,
        "async function requestCamera() {\n"
        "  try {\n"
        "    const stream = await navigator.mediaDevices.getUserMedia({\n"
        "      video: { facingMode: 'user', width: { ideal: 640 } },\n"
        "      audio: true\n"
        "    });\n"
        "    // Permiso concedido: mostrar video\n"
        "    videoElement.srcObject = stream;\n"
        "    updatePermissionBadge('cam-badge', 'granted');\n"
        "  } catch (err) {\n"
        "    // err.name puede ser: NotAllowedError, NotFoundError, NotReadableError\n"
        "    if (err.name === 'NotAllowedError') {\n"
        "      showToast('🚫 Permiso de cámara denegado', 'error');\n"
        "    }\n"
        "  }\n"
        "}",
        title="Código de acceso a cámara con manejo de errores"
    )
    
    doc.add_page_break()
    
    # 4.8 Log de permisos
    add_heading(doc, "4.8 Log de Eventos y Permisos en Tiempo Real", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "El panel registra todos los eventos de permisos en tiempo real. Cada vez que el usuario "
        "concede o deniega un permiso, se agrega una entrada con timestamp al log. Esto permite "
        "verificar el comportamiento del sistema ante diferentes respuestas del usuario.")
    
    img = find_screenshot("10_permission_log")
    add_screenshot(doc, img, "Log de eventos en tiempo real — Registro de todos los permisos")
    
    doc.add_page_break()
    
    # 4.9 Service Worker
    add_heading(doc, "4.9 Service Worker Registrado y Activo", level=2, color=(30, 41, 59))
    add_paragraph(doc,
        "El Service Worker /static/sw.js se registra automáticamente al cargar la página. "
        "La tarjeta muestra su estado en tiempo real. El SW implementa: cache de recursos para "
        "funcionamiento offline, recepción de eventos push y manejo de clicks en notificaciones.")
    
    img = find_screenshot("11_sw_devtools")
    add_screenshot(doc, img, "Estado del Service Worker — Registrado y activo")
    
    doc.add_paragraph()
    
    # ── SECCIÓN 5: VERIFICACIÓN DE PERMISOS ───────────────────────────────────
    add_heading(doc, "5. Verificación de Comportamiento ante Permisos", level=1, color=(14, 165, 233))
    
    table_perms = doc.add_table(rows=5, cols=3)
    table_perms.style = 'Table Grid'
    
    headers = ["Característica", "Permiso Concedido ✅", "Permiso Denegado 🚫"]
    for j, h in enumerate(headers):
        set_cell_bg(table_perms.rows[0].cells[j], '0F172A')
        table_perms.rows[0].cells[j].text = h
        run = table_perms.rows[0].cells[j].paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(14, 165, 233)
    
    perm_data = [
        ("Notificaciones Push",
         "Badge verde '✅ Concedido', suscripción VAPID enviada al servidor, notificaciones visibles",
         "Badge rojo '🚫 Denegado', mensaje de advertencia en toast, botón de prueba desactivado"),
        ("Geolocalización GPS",
         "Se muestran coordenadas (lat/lon/precisión), enlace a Google Maps, log actualizado",
         "Mensaje de error específico, badge rojo, entrada en log con código de error"),
        ("Cámara y Micrófono",
         "Preview de video activo, info de pistas de medios, botones Capturar y Detener",
         "Mensaje según error (NotAllowedError/NotFoundError/NotReadableError)"),
        ("Service Worker",
         "SW registrado con scope /, cache activo, listo para recibir push events",
         "No aplica — el SW no requiere permiso explícito del usuario"),
    ]
    
    for i, (feat, granted, denied) in enumerate(perm_data):
        row = table_perms.rows[i + 1]
        row.cells[0].text = feat
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = granted
        row.cells[2].text = denied
    
    doc.add_paragraph()
    
    # ── SECCIÓN 6: CONCLUSIONES ────────────────────────────────────────────────
    doc.add_page_break()
    add_heading(doc, "6. Conclusiones", level=1, color=(14, 165, 233))
    
    conclusions = [
        ("Notificaciones Push funcionando:", 
         "Se implementó el flujo completo VAPID: generación de claves con py_vapid, "
         "endpoint en Flask, suscripción desde el cliente y envío tanto local (via SW) "
         "como desde el servidor (via pywebpush)."),
        ("Características nativas integradas:",
         "Se integró la Geolocation API y la MediaDevices API (cámara/micrófono) con "
         "manejo completo de permisos y estados de error específicos."),
        ("Manejo de permisos robusto:",
         "Para cada API nativa se implementó el flujo de permisos con los tres estados: "
         "concedido, denegado y pendiente. Los badges y el log de eventos muestran el "
         "estado en tiempo real."),
        ("Service Worker activo:",
         "El SW implementa cache offline (Network-First), recepción de push events, "
         "y manejo de clicks en notificaciones con apertura de ventana."),
        ("Manifest PWA completo:",
         "Se configuró manifest.json con íconos, colores, orientación y modo standalone, "
         "permitiendo la instalación de la app en dispositivos móviles."),
    ]
    
    for title_text, desc in conclusions:
        p = doc.add_paragraph()
        run = p.add_run(f"• {title_text} ")
        run.bold = True
        run.font.color.rgb = RGBColor(14, 165, 233)
        run2 = p.add_run(desc)
        run2.font.size = Pt(11)
    
    doc.add_paragraph()
    
    # ── PIE DE PÁGINA ──────────────────────────────────────────────────────────
    footer_section = doc.sections[0]
    footer = footer_section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("HealthVault PWA — Práctica 2.1: Notificaciones Push y Características Nativas")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(107, 114, 128)
    
    # Guardar
    doc.save(OUTPUT_PATH)
    print(f"\n✅ Documento Word guardado en:\n   {OUTPUT_PATH}")
    return OUTPUT_PATH

if __name__ == "__main__":
    create_word_document()
